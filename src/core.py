import csv, json, re
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
from jinja2 import Environment, FileSystemLoader, StrictUndefined
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib import colors
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES = ROOT / "templates"
DATA = ROOT / "data"
OUTPUT = ROOT / "output"
COUNTERS = DATA / "series.json"
for folder in (TEMPLATES, DATA, OUTPUT):
    folder.mkdir(parents=True, exist_ok=True)

TYPES = {
    "Attestation": ("attestation", "AT"),
    "Convocation": ("convocation", "CV"),
    "Reçu": ("recu", "RC"),
    "Bulletin": ("bulletin", "BL"),
}

def read_rows(path):
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            return list(csv.DictReader(stream))
    if path.suffix.lower() == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        values = ws.iter_rows(values_only=True)
        headers = next(values, [])
        headers = [str(x).strip().lower() if x is not None else "" for x in headers]
        rows = []
        for row in values:
            rows.append({headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]})
        wb.close()
        return rows
    raise ValueError("Choisissez un fichier .csv ou .xlsx.")

class DocumentEngine:
    def __init__(self):
        self.env = Environment(loader=FileSystemLoader(str(TEMPLATES)), autoescape=False,
                               undefined=StrictUndefined, trim_blocks=True, lstrip_blocks=True)
        if not COUNTERS.exists():
            COUNTERS.write_text(json.dumps({key: 0 for key, _ in TYPES.values()}, indent=2), encoding="utf-8")

    def next_id(self, kind):
        key, prefix = TYPES[kind]
        try:
            counters = json.loads(COUNTERS.read_text(encoding="utf-8"))
        except Exception:
            counters = {}
        counters[key] = int(counters.get(key, 0)) + 1
        COUNTERS.write_text(json.dumps(counters, indent=2, ensure_ascii=False), encoding="utf-8")
        return f"{prefix}-{datetime.now().year}-{counters[key]:04d}"

    def render(self, kind, data):
        return self.env.get_template(f"{TYPES[kind][0]}.txt").render(**data)

    def generate(self, kind, data, formats):
        name = str(data.get("nom") or "").strip()
        first = str(data.get("prenom") or "").strip()
        if not name or not first:
            raise ValueError("Le nom et le prénom sont obligatoires.")
        normalized = {
            "civilite": str(data.get("civilite") or "M.").strip(),
            "nom": name, "prenom": first,
            "date": str(data.get("date") or datetime.now().strftime("%d/%m/%Y")),
            "motif": str(data.get("motif") or "Non précisé"),
            "montant": str(data.get("montant") or "0"),
            "note": str(data.get("note") or "0"),
            "email": str(data.get("email") or ""),
        }
        number = self.next_id(kind)
        normalized["numero"] = number
        text = self.render(kind, normalized)
        safe = re.sub(r"[^A-Za-z0-9_-]+", "_", f"{name}_{first}").strip("_")
        stem = OUTPUT / f"{number}_{safe}"
        paths = []
        for fmt in formats:
            path = stem.with_suffix(".pdf" if fmt == "PDF" else ".docx")
            if fmt == "PDF":
                self._pdf(text, path)
            elif fmt == "Word":
                self._word(text, path)
            paths.append(path)
        return number, paths

    @staticmethod
    def _pdf(text, path):
        doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=24*mm, rightMargin=24*mm,
                                topMargin=24*mm, bottomMargin=22*mm)
        styles = getSampleStyleSheet()
        title = ParagraphStyle("TitleCustom", parent=styles["Title"], alignment=TA_CENTER,
                               fontSize=16, leading=21, textColor=colors.HexColor("#234B52"),
                               spaceAfter=16)
        body = ParagraphStyle("BodyCustom", parent=styles["BodyText"], fontSize=11,
                              leading=17, spaceAfter=9)
        lines = text.splitlines()
        story = [Paragraph(escape(lines[0] if lines else "DOCUMENT"), title),
                 HRFlowable(width="100%", thickness=1, color=colors.HexColor("#6A9C91")),
                 Spacer(1, 16)]
        story.extend(Paragraph(escape(line) if line.strip() else "&nbsp;", body) for line in lines[1:])
        doc.build(story)

    @staticmethod
    def _word(text, path):
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(.8)
        section.bottom_margin = Inches(.8)
        for i, line in enumerate(text.splitlines()):
            paragraph = doc.add_paragraph()
            if i == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(line)
                run.bold = True
                run.font.size = Pt(16)
            else:
                paragraph.add_run(line)
        doc.save(path)

    def read_template(self, kind):
        return (TEMPLATES / f"{TYPES[kind][0]}.txt").read_text(encoding="utf-8")

    def save_template(self, kind, text):
        (TEMPLATES / f"{TYPES[kind][0]}.txt").write_text(text, encoding="utf-8")
